import hashlib
from io import BytesIO
import os
import re
from uuid import uuid4

from ..domain.errors import ScholarError, not_found
from ..domain.evidence import DocumentVersion
from ..domain.operations import ExportJob, ExportArtifact
from ..domain.reports import DraftAcceptance, VerificationRecord
from ..persistence.records import RecordStore
from ..persistence.report_store import ReportStore
from ..persistence.run_store import fingerprint
from .citations import render_citations

FORMATS = {'markdown': ('md', 'text/markdown'), 'docx': ('docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'), 'bibtex': ('bib', 'application/x-bibtex')}


class ExportService:
    def __init__(self, db, settings):
        self.db = db
        self.settings = settings
        self.records = RecordStore(db)
        self.reports = ReportStore(db)
        self.root = (settings.data_dir / 'exports').resolve()

    def get(self, job_id, conn=None):
        return self.records.get('export', job_id, ExportJob, conn)

    def list(self, version_id):
        return [j for j in self.records.list('export', ExportJob) if j.report_version_id == version_id]

    def request(self, version_id, format, verification_id, draft_acceptance_id, operation_key):
        if format not in FORMATS:
            raise ScholarError(422, 'invalid_export_format', 'Choose markdown, docx or bibtex')
        inputs = {'format': format, 'verification_id': verification_id, 'draft_acceptance_id': draft_acceptance_id}
        scope, digest = f'export:{version_id}', fingerprint(inputs)
        with self.db.transaction() as conn:
            previous = conn.execute('SELECT * FROM idempotency WHERE scope=? AND operation_key=?', (scope, operation_key)).fetchone()
            if previous:
                if previous['request_hash'] != digest:
                    raise ScholarError(409, 'idempotency_conflict', 'Export key has different inputs')
                return self.get(previous['entity_id'], conn)
            report = self.reports.get(version_id, conn)
            check = self.records.get('verification', verification_id, VerificationRecord, conn)
            if check.report_version_id != report.id:
                raise ScholarError(409, 'verification_version_mismatch', 'Export must use a check of this exact version')
            draft = check.quality_status != 'passed'
            if draft:
                if not draft_acceptance_id:
                    raise ScholarError(409, 'draft_acceptance_required', 'Accept this exact issue-bearing draft before exporting it')
                acceptance = self.records.get('acceptance', draft_acceptance_id, DraftAcceptance, conn)
                if acceptance.report_version_id != report.id or acceptance.verification_record_id != check.id:
                    raise ScholarError(409, 'draft_acceptance_mismatch', 'Draft acceptance does not match this report and check')
            job = ExportJob(project_id=report.project_id, report_version_id=report.id, format=format,
                verification_id=check.id, draft_acceptance_id=draft_acceptance_id, draft=draft,
                document_version_ids=report.manifest.get('document_version_ids', []))
            self.records.insert('export', job, project_id=report.project_id, links=[('report', report.id), ('verification', check.id)], conn=conn)
            conn.execute('INSERT INTO idempotency(scope,operation_key,request_hash,entity_id) VALUES (?,?,?,?)', (scope, operation_key, digest, job.id))
            return job

    def resolve(self, key, format):
        if format not in FORMATS or not re.fullmatch(r'[0-9a-f]{64}', key or ''):
            raise ScholarError(422, 'invalid_export_path', 'Invalid export identifier')
        path = (self.root / f'{key}.{FORMATS[format][0]}').resolve()
        if not path.is_relative_to(self.root) or not path.is_file():
            raise not_found('export file')
        return path

    def build(self, job_id):
        job = self.get(job_id)
        report = self.reports.get(job.report_version_id)
        rendered = render_citations(report)
        marker = 'DRAFT — issues explicitly accepted by the user; model checks are not factual guarantees.' if job.draft else 'Model support checks recorded; human scholarly review remains necessary.'
        audit = f'Report version: {report.id}\nVerification: {job.verification_id}\n{marker}'
        source_lines = []
        for did in job.document_version_ids:
            document = self.records.get('document', did, DocumentVersion)
            source_lines.append(f'{did} | {document.kind} | SHA256 {document.sha256}')
        appendix = '\n\n## Source versions\n\n' + '\n\n'.join(source_lines)
        coverage = report.manifest.get('coverage', {})
        if coverage:
            appendix += '\n\n## Reading coverage\n\n' + '\n\n'.join(
                f"{label}: {', '.join(coverage.get(key, [])) or 'None'}" for label, key in [
                    ('Included', 'selected_paper_version_ids'), ('Abstract only', 'abstract_only_paper_version_ids'), ('Unreadable', 'unreadable_paper_version_ids')])
        markdown = audit + '\n\n' + rendered.markdown + appendix
        if job.format == 'markdown':
            data = markdown.encode('utf-8')
        elif job.format == 'docx':
            from docx import Document
            from docx.shared import Pt
            document = Document()
            document.styles['Normal'].font.size = Pt(10.5)
            for paragraph in markdown.split('\n\n'):
                if paragraph.startswith('## '):
                    document.add_heading(paragraph[3:], level=2)
                elif paragraph.startswith('# '):
                    document.add_heading(paragraph[2:], level=1)
                else:
                    document.add_paragraph(paragraph)
            buffer = BytesIO()
            document.save(buffer)
            data = buffer.getvalue()
        else:
            def escape(value):
                replacements = {'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}', '\n': ' '}
                return ''.join(replacements.get(char, char) for char in str(value))
            entries = ['% ' + audit.replace('\n', '\n% ')]
            for ref in rendered.references:
                fields = {'title': ref.get('title'), 'author': ' and '.join(ref.get('authors') or []), 'year': ref.get('year'),
                          'doi': ref.get('identifiers', {}).get('doi'), 'eprint': ref.get('identifiers', {}).get('arxiv')}
                values = ',\n'.join(f'  {key} = {{{escape(value)}}}' for key, value in fields.items() if value)
                entries.append('@article{' + ref['key'] + ',\n' + values + '\n}')
            data = '\n\n'.join(entries).encode('utf-8')
        digest = hashlib.sha256(data).hexdigest()
        self.root.mkdir(parents=True, exist_ok=True)
        destination = self.root / f'{digest}.{FORMATS[job.format][0]}'
        temporary = self.root / (str(uuid4()) + '.tmp')
        try:
            temporary.write_bytes(data)
            os.replace(temporary, destination)
        finally:
            temporary.unlink(missing_ok=True)
        return ExportArtifact(report_version_id=report.id, storage_key=digest, format=job.format, size=len(data))
